import io
import json
import os
import re
from collections import Counter
from difflib import SequenceMatcher
from flask import Flask, request, render_template, send_file, jsonify
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 * 1024

MASKED_WORDS = [
    ("k** kh** q**n", "kéo khóa quần"),
    ("c*n m** d***", "cắn môi dưới"),
    ("k*ch th*ch", "kích thích"),
    ("ph*ng đ*ng", "phóng đãng"),
    ("b**n th**", "biến thái"),
    ("th*n d***", "thân dưới"),
    ("tr*n tr**", "trần trụi"),
    ("kh*u gi*o", "khẩu giao"),
    ("kh*** c*m", "khoái cảm"),
    ("g**t ch*t", "giết chết"),
    ("d*m đ*ng", "dâm đãng"),
    ("b*** tối", "buổi tối"),
    ("kh*ng b*", "khủng bố"),
    ("l*m t*nh", "làm tình"),
    ("l**m m*t", "liếm mút"),
    ("d*c v*ng", "dục vọng"),
    ("h*m m**n", "ham muốn"),
    ("l*n đ*nh", "lên đỉnh"),
    ("b*p ch*t", "bóp chết"),
    ("th* d**c", "thở dốc"),
    ("th* d*c", "thở dốc"),
    ("v**t v*", "vuốt ve"),
    ("l* m*ng", "lỗ mãng"),
    ("k** r*n", "kêu rên"),
    ("c** nh*", "cậu nhỏ"),
    ("x** n*n", "xoa nắn"),
    ("c*m v**", "cắm vào"),
    ("n*n b*p", "nắn bóp"),
    ("r*n r*", "rên rỉ"),
    ("v*n v*", "vân vê"),
    ("v* v*n", "ve vãn"),
    ("b*n r*", "bắn ra"),
    ("b* q**", "ba que"),
    ("m* t**", "ma túy"),
    ("ch*ch", "chịch"),
    ("h**p", "hiếp"),
    ("l**m", "liếm"),
    ("g**t", "giết"),
    ("ch*t", "chết"),
    ("d*m", "dâm"),
    ("b*n", "bắn"),
    ("m*t", "mút"),
]

REMOVE_LINES = [
    re.compile(r"^\s*Nguồn:\s*(?:\[read\.st\]\(https?://(?:www\.)?read\.st/?\)|https?://(?:www\.)?read\.st/?|read\.st)\s*$", re.I),
    re.compile(r"^\s*-{2,}\s*oOo\s*-{2,}\s*$", re.I),
]

WORD_FIXES = [
    ("th* d**c → thở dốc", re.compile(r"(?<!\w)th\*\s+d\*\*c(?!\w)", re.I), "thở dốc"),
    ("r*n r* → rên rỉ", re.compile(r"(?<!\w)r\*n\s+r\*(?!\w)", re.I), "rên rỉ"),
]
PATTERNS = [
    ("Ký hiệu ^p lặp", re.compile(r"(?:\^p){2,}", re.I), "^p"),
    ("Khoảng trắng lặp", re.compile(r"(?<=\S)[ \t]{2,}"), " "),
    ("Dấu cách trước dấu câu", re.compile(r"[ \t]+(?=[,:;!?])"), ""),
    ("Dấu cách trước dấu chấm", re.compile(r"[ \t]+(?=\.(?!\.))"), ""),
    ("Thiếu khoảng trắng sau dấu phẩy", re.compile(r',(?=[^\s\d"”’.,])'), ", "),
    ("Ký hiệu ^+^+ dính chữ", re.compile(r"(?<=[^\W\d_])(?=\^\+\^\+)"), " "),
]

COMPILED_MASKED_WORDS = [(re.compile(r'(?<!\w)' + re.escape(src) + r'(?!\w)', re.I), dst) for src, dst in MASKED_WORDS]

def fix_text(text, counts):
    if '*' in text:
        for pattern, target in COMPILED_MASKED_WORDS:
            text, changed = pattern.subn(lambda m: target.capitalize() if m.group(0)[:1].isupper() else target, text)
            counts['Khôi phục từ bị che'] += changed
    for name, pattern, replacement in PATTERNS:
        text, changed = pattern.subn(replacement, text)
        counts[name] += changed
    return text

def replace_runs(paragraph, old, new):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new)
        return
    if len(runs) == 1:
        runs[0].text = new
        return
    chars = [(i, c) for i, run in enumerate(runs) for c in run.text]
    if len(chars) != len(old):
        return False
    for op, a, b, c, d in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if op != "equal":
            target = chars[a][0] if a < len(chars) else len(runs)-1
            chars[a:b] = [(target, ch) for ch in new[c:d]]
    pieces = [""] * len(runs)
    for i, ch in chars:
        pieces[i] += ch
    for run, value in zip(runs, pieces):
        run.text = value
    return True

def remove_blank_paragraphs(doc):
    removed = 0
    for paragraph in list(doc.paragraphs):
        if not paragraph.text.strip() and not paragraph._p.xpath(".//w:drawing | .//w:br | .//w:object | .//w:pict"):
            el = paragraph._element
            el.getparent().remove(el)
            removed += 1
    return removed

def is_heading(text):
    title = text.strip()
    return bool(re.match(r"^(?:Chương|Chapter)\s+\d+(?:\s*[:：.\-–].*)?$", title, re.I)) or title.casefold() in ("văn án", "giới thiệu", "mục lục")

def iter_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph

def clean(data, options):
    doc = Document(io.BytesIO(data))
    counts = Counter()
    if options.get("narrow"):
        for section in doc.sections:
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(1.27)
    for paragraph in iter_paragraphs(doc):
        original = paragraph.text
        if ('read.st' in original.lower() or 'ooo' in original.lower()) and any(pattern.fullmatch(original.strip()) for pattern in REMOVE_LINES):
            element = paragraph._element
            element.getparent().remove(element)
            counts['Xóa dòng nguồn / phân cách'] += 1
            continue
        local = Counter()
        cleaned = fix_text(original, local)
        if cleaned != original:
            if replace_runs(paragraph, original, cleaned) is not False:
                counts.update(local)
            else:
                counts["Đoạn cần kiểm tra"] += 1
        if options.get("headings") and is_heading(paragraph.text):
            paragraph.style = "Heading 2"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = "Tahoma"
                run.font.size = Pt(11)
            counts["Chuẩn hóa Heading 2"] += 1
    counts["Đoạn trống đã xóa"] = remove_blank_paragraphs(doc)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output, dict(counts)

@app.get("/")
def index():
    return render_template("index.html")

@app.post("/api/clean")
def api_clean():
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".docx"):
        return jsonify(error="Vui lòng tải file .docx"), 400
    options = {
        "narrow": request.form.get("narrow") == "true",
        "headings": request.form.get("headings") == "true"
    }
    try:
        output, counts = clean(upload.read(), options)
    except Exception:
        app.logger.exception('Document processing error')
        return jsonify(error='Không xử lý được tài liệu. Vui lòng thử một file DOCX nhỏ hơn.'), 422
    original_filename = os.path.splitext(secure_filename(upload.filename))[0] or "document"
    response = send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=original_filename + "_da_sua.docx",
    )
    response.headers["X-Corrections"] = json.dumps(counts, ensure_ascii=True)
    response.headers["Access-Control-Expose-Headers"] = "X-Corrections"
    return response

@app.get("/health")
def health():
    return {"ok": True}

if __name__ == "__main__":
    app.run(debug=True)
